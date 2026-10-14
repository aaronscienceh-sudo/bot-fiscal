"""Parser de CSV/TSV e Excel.

Detecta automaticamente se a planilha é um EXTRATO (transações) ou uma
lista de FATURAS, mapeando colunas por aliases de cabeçalho (PT/EN).
"""
from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

from ..domain.models import Invoice, Transaction
from ..domain.rules import extract_reference, normalize_name, parse_amount, parse_date

# aliases -> campo
TXN_COLUMNS = {
    "date": {"data", "date", "dt", "dt_lancamento", "lancamento", "data_lancamento", "dia"},
    "amount": {"valor", "amount", "valor_r", "debito", "credito", "vlr", "montante"},
    "description": {"descricao", "historico", "historico_agencia", "descri", "description",
                    "informacao", "info", "detalhe", "lancamento"},
    "reference": {"referencia", "numero", "número", "num", "doc", "protocolo", "nf", "boleto"},
    "counterparty": {"favorecido", "favorecida", "pagador", "contraparte", "nome",
                     "fornecedor", "beneficiario", "historico"},
    "account": {"conta", "agencia", "agc", "agencia_cuenta"},
}

INV_COLUMNS = {
    "number": {"numero", "número", "num", "nf", "nfe", "fatura", "nota", "nota_fiscal", "id"},
    "issuer": {"emissor", "fornecedor", "empresa", "issuer", "razao_social", "cnpj", "cliente"},
    "issue_date": {"emissao", "emissão", "data_emissao", "data_emissao", "dt_emissao"},
    "due_date": {"vencimento", "venc", "due", "vencimento", "data_vencimento", "pagamento"},
    "amount": {"valor", "amount", "total", "vlr", "valor_total", "montante"},
    "description": {"descricao", "descri", "description", "item", "produto", "servico"},
}


def _match_header(header: str, aliases: dict[str, set[str]]) -> str | None:
    h = normalize_name(header)
    for field, names in aliases.items():
        if h in names:
            return field
    # correspondência parcial (ex.: "valor (r$)")
    for field, names in aliases.items():
        for n in names:
            if n in h or h in n:
                return field
    return None


def _detect_columns(headers: list[str], aliases: dict[str, set[str]]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, h in enumerate(headers):
        field = _match_header(h, aliases)
        if field and field not in mapping:
            mapping[field] = idx
    return mapping


def _is_transactions(cols: dict[str, int]) -> bool:
    # extrato tem data+valor+descrição; fatura tem número+emissor
    has_txn = {"date", "amount"} <= set(cols) and "description" in cols
    has_inv = {"number", "amount"} <= set(cols) and "issuer" in cols
    return has_txn or (has_txn and not has_inv)


def _row_to_transaction(row: list[str], cols: dict[str, int], idx: int, source: str) -> Transaction | None:
    def get(field: str) -> str:
        i = cols.get(field)
        return row[i].strip() if i is not None and i < len(row) else ""

    amount_raw = get("amount")
    date_raw = get("date")
    if not amount_raw and not date_raw:
        return None
    amount = parse_amount(amount_raw)
    d = parse_date(date_raw)
    description = get("description")
    ref = get("reference") or extract_reference(description)
    h = hashlib.md5(f"{source}|{d}|{amount}|{description}".encode()).hexdigest()[:8]
    return Transaction(
        id=f"TXN-{h}",
        date=d,
        amount=amount,
        description=description,
        reference=ref or None,
        counterparty=get("counterparty") or None,
        account=get("account") or None,
        source_file=source,
    )


def _row_to_invoice(row: list[str], cols: dict[str, int], idx: int, source: str) -> Invoice | None:
    def get(field: str) -> str:
        i = cols.get(field)
        return row[i].strip() if i is not None and i < len(row) else ""

    number = get("number")
    amount_raw = get("amount")
    if not number and not amount_raw:
        return None
    h = hashlib.md5(f"{source}|{number}".encode()).hexdigest()[:8]
    return Invoice(
        id=f"INV-{h}",
        number=number or f"GEN-{idx}",
        issuer=get("issuer") or "Desconhecido",
        issue_date=parse_date(get("issue_date")),
        due_date=parse_date(get("due_date")),
        amount=parse_amount(amount_raw),
        description=get("description") or None,
        source_file=source,
    )


def _process_rows(rows: list[list[str]], source: str) -> tuple[list[Transaction], list[Invoice], list[str]]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        return [], [], ["Arquivo vazio ou sem linhas de dados."]
    headers = [str(h) for h in rows[0]]
    data = rows[1:]

    cols_txn = _detect_columns(headers, TXN_COLUMNS)
    cols_inv = _detect_columns(headers, INV_COLUMNS)
    kind = "transactions" if _is_transactions(cols_txn) else "invoices"

    txns: list[Transaction] = []
    invs: list[Invoice] = []
    errors: list[str] = []

    if kind == "transactions":
        for i, row in enumerate(data, start=2):
            t = _row_to_transaction(row, cols_txn, i, source)
            if t:
                txns.append(t)
    else:
        for i, row in enumerate(data, start=2):
            inv = _row_to_invoice(row, cols_inv, i, source)
            if inv:
                invs.append(inv)

    if not txns and not invs:
        errors.append("Nenhuma linha reconhecida. Verifique os cabeçalhos.")
    return txns, invs, errors


def parse_csv(path: str | Path, delimiter: str | None = None) -> tuple[list[Transaction], list[Invoice], list[str]]:
    path = Path(path)
    source = path.name
    with path.open(newline="", encoding="utf-8-sig") as f:
        # detecta o delimitador
        sample = f.read(4096)
        f.seek(0)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = "," if "," in sample else ";"
        reader = csv.reader(f, delimiter=delimiter)
        rows = [row for row in reader]
    return _process_rows(rows, source)


def parse_excel(path: str | Path) -> tuple[list[Transaction], list[Invoice], list[str]]:
    from openpyxl import load_workbook
    path = Path(path)
    source = path.name
    wb = load_workbook(path, data_only=True, read_only=True)
    all_txns: list[Transaction] = []
    all_invs: list[Invoice] = []
    errors: list[str] = []
    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
        txns, invs, err = _process_rows(rows, f"{source}::{ws.title}")
        all_txns.extend(txns)
        all_invs.extend(invs)
        errors.extend(err)
    wb.close()
    return all_txns, all_invs, errors
